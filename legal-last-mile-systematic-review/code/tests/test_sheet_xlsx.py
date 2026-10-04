#!/usr/bin/env python3
"""Tests for code/analysis/sheet_xlsx.py: the drop-down .xlsx versions of the A16 sheet and the second-extractor sheet hold exactly the CSV's content, carry the
drop-downs, and a FILLED workbook is read correctly by apply_a16_adjudications.py and score_second_extractor_sheet.py. Uses temp copies; the committed workbooks are
only read. Run from legal-last-mile-systematic-review/:  python3 code/tests/test_sheet_xlsx.py"""
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'code/analysis'))
sys.path.insert(0, str(ROOT / 'code/screening'))
import apply_a16_adjudications as apply_mod  # noqa: E402
import score_second_extractor_sheet as scorer  # noqa: E402
import sheet_xlsx as sx  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402


class TestSheets(unittest.TestCase):
    def test_committed_workbooks_match_their_csv(self):
        for p in sx.SPECS:
            self.assertTrue(sx.in_sync(p), f'{sx.xlsx_path(p).name} is out of date: run python3 code/analysis/sheet_xlsx.py')

    def test_drop_downs_and_header(self):
        for p, (dropdowns, _) in sx.SPECS.items():
            ws = load_workbook(sx.xlsx_path(p)).worksheets[0]
            self.assertEqual(ws.freeze_panes, 'A2')
            got = {str(dv.sqref).split(':')[0].rstrip('0123456789'): dv.formula1 for dv in ws.data_validations.dataValidation}
            header = [c.value for c in ws[1]]
            self.assertEqual(len(got), len(dropdowns))
            for h, allowed in dropdowns.items():
                col = get_column_letter(header.index(h) + 1)
                self.assertEqual(got[col], '"' + ','.join(allowed) + '"')

    def test_ids_stay_text_and_empty_rows_are_skipped(self):
        with tempfile.TemporaryDirectory() as d:
            from openpyxl import Workbook
            wb = Workbook(); ws = wb.active
            ws.append(['record_id', 'n']); ws.append(['R0123456E789', 12.0]); ws.append([None, None]); ws.append(['S001', 'x'])
            f = Path(d) / 't.xlsx'; wb.save(f)
            h, rows = sx.read_xlsx(f)
            self.assertEqual(h, ['record_id', 'n']); self.assertEqual([r['n'] for r in rows], ['12', 'x']); self.assertEqual(len(rows), 2)

    def test_filled_a16_workbook_is_read_by_the_apply_script(self):
        with tempfile.TemporaryDirectory() as d:
            f = Path(d) / 'filled.xlsx'
            sx.write_xlsx(sx.A16_CSV, f)
            wb = load_workbook(f); ws = wb.worksheets[0]
            header = [c.value for c in ws[1]]
            rid = header.index('record_id') + 1; dec = header.index(apply_mod.DEC_COL) + 1; code = header.index(apply_mod.CODE_COL) + 1
            ws.cell(row=2, column=dec, value='exclude'); ws.cell(row=2, column=code, value='e3')
            ws.cell(row=3, column=dec, value='include')
            first, second = ws.cell(row=2, column=rid).value, ws.cell(row=3, column=rid).value
            wb.save(f)
            rows = apply_mod.read_sheet(f)
            by = {r['record_id']: r for r in rows}
            self.assertEqual(by[first.upper()][apply_mod.DEC_COL], 'exclude'); self.assertEqual(by[second.upper()][apply_mod.DEC_COL], 'include')
            self.assertEqual(apply_mod.norm_code(by[first.upper()][apply_mod.CODE_COL]), 'E03')
            self.assertEqual(sum(1 for r in rows if r[apply_mod.DEC_COL]), 2)

    def test_filled_second_extractor_workbook_is_scored(self):
        with tempfile.TemporaryDirectory() as d:
            f = Path(d) / 'filled.xlsx'
            sx.write_xlsx(sx.SE_CSV, f)
            wb = load_workbook(f); ws = wb.worksheets[0]
            col = [c.value for c in ws[1]].index(scorer.VERDICT) + 1
            for r, v in ((2, 'Y'), (3, 'N'), (4, 'cannot_tell')):
                ws.cell(row=r, column=col, value=v)
            wb.save(f)
            R = scorer.score(scorer.read_rows(str(f)))
            self.assertEqual((sum(R['tot'].values()), sum(R['bad'].values())), (2, 1)); self.assertEqual(R['invalid'], [])

    def test_awkward_workbooks(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            # padded header, a second sheet, capitalised values, an upper-case extension, trailing blank rows
            wb = Workbook(); ws = wb.active
            ws.append([' record_id ', apply_mod.DEC_COL + ' ', apply_mod.CODE_COL]); ws.append([' r0abc ', 'Include', None]); ws.append([None, None, None]); ws.append([None, None, None])
            wb.create_sheet('notes').append(['ignored'])
            f = d / 'X.XLSX'; wb.save(f)
            rows = apply_mod.read_sheet(f)
            self.assertEqual(len(rows), 1); self.assertEqual(rows[0]['record_id'], 'R0ABC'); self.assertEqual(rows[0][apply_mod.DEC_COL], 'Include')
            plan, errors = apply_mod.plan(rows, {'R0ABC': dict(final_decision='exclude', full_text_decision='exclude', exclusion_reason='E01', notes='')})
            self.assertEqual(errors, []); self.assertEqual(plan[0][0], 'reverse')  # 'Include' is read case-insensitively
            # missing required column
            wb = Workbook(); wb.active.append(['record_id', 'something']); g = d / 'm.xlsx'; wb.save(g)
            with self.assertRaises(ValueError):
                apply_mod.read_sheet(g)
            with self.assertRaises(SystemExit):
                scorer.read_rows(str(g))
            # not a workbook at all
            h = d / 'broken.xlsx'; h.write_bytes(b'this is not a zip file')
            with self.assertRaises(ValueError) as cm:
                apply_mod.read_sheet(h)
            self.assertIn('could not read', str(cm.exception))
            with self.assertRaises(SystemExit) as cm2:
                scorer.read_rows(str(h))
            self.assertIn('could not read', str(cm2.exception))


if __name__ == '__main__':
    unittest.main(verbosity=1)
