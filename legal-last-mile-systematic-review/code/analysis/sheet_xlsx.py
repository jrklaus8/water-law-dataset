#!/usr/bin/env python3
"""Spreadsheet (.xlsx) versions of the two sheets the researcher must fill in, and a reader for them.

A CSV opened in Excel or Google Sheets invites typos in the decision columns and mangles ids. The .xlsx versions made here hold exactly the CSV's rows and columns
(every cell is text, so nothing is reformatted), freeze the header, wrap long text, and put a drop-down on each column the researcher fills in (A16: include/exclude and
E01-E12; second extractor: Y / N / cannot_tell), so an invalid entry is refused as it is typed. `read_xlsx(path)` returns (header, rows) in the same shape the CSV
readers use, so `apply_a16_adjudications.py` and `score_second_extractor_sheet.py` accept a filled .xlsx directly (no "save as CSV" step).
Usage (from legal-last-mile-systematic-review/):  python3 code/analysis/sheet_xlsx.py            # (re)write the two .xlsx files next to their CSVs
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

csv.field_size_limit(sys.maxsize)
ROOT = Path(__file__).resolve().parents[2]
A16_CSV = ROOT / '02_screening/full_text/A16_ADJUDICATION_SHEET_2026-10-04.csv'
SE_CSV = ROOT / '03_extraction/second_extractor/second_extractor_sheet_BLANK_2026-10-04.csv'
# (csv, columns that get a drop-down: header -> allowed values, columns shown wide)
SPECS = {
    A16_CSV: ({'researcher_decision (include/exclude)': ['include', 'exclude'],
               'researcher_exclusion_code (E01-E12)': [f'E{i:02d}' for i in range(1, 13)]},
              {'title': 50, 'ai_reasoning': 60, 'codex_rationale': 60, 'gemini_rationale': 60, 'claude_triage_note': 50, 'researcher_comment': 40}),
    SE_CSV: ({'agrees (Y/N/cannot_tell)': ['Y', 'N', 'cannot_tell']},
             {'citation': 50, 'what_to_check': 45, 'ai_value': 45, 'second_extractor_value': 45, 'where_the_AI_looked': 40, 'comment': 35}),
}


def read_csv(path: Path):
    with open(path, encoding='utf-8', newline='') as f:
        rd = csv.reader(f)
        header = next(rd)
        return header, [r for r in rd]


def xlsx_path(csv_path: Path) -> Path:
    return csv_path.with_suffix('.xlsx')


def write_xlsx(csv_path: Path, out: Path | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    header, rows = read_csv(csv_path)
    dropdowns, wide = SPECS[csv_path]
    wb = Workbook(); ws = wb.active; ws.title = 'sheet'
    ws.append(header)
    for r in rows:
        ws.append(r + [''] * (len(header) - len(r)))
    for row in ws.iter_rows():
        for c in row:
            c.number_format = '@'  # text: ids such as S001 or R0532... and DOIs are never reinterpreted
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    for i, h in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = wide.get(h, min(28, max(10, len(h) + 2)))
    for h, allowed in dropdowns.items():
        col = get_column_letter(header.index(h) + 1)
        dv = DataValidation(type='list', formula1='"' + ','.join(allowed) + '"', allow_blank=True, showErrorMessage=True,
                            errorTitle='Not allowed', error='Choose one of: ' + ', '.join(allowed))
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{len(rows) + 1}')
    out = out or xlsx_path(csv_path)
    wb.save(out)
    return out


def read_xlsx(path):
    """(header, rows) from the first worksheet; rows are dicts with '' for empty cells; whole-number floats lose their '.0'; fully blank rows are skipped."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else '' for h in next(it)]
    def cell(v):
        if v is None:
            return ''
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)
    rows = []
    for r in it:
        vals = [cell(v) for v in r]
        if any(x.strip() for x in vals):
            rows.append(dict(zip(header, vals + [''] * (len(header) - len(vals)))))
    wb.close()
    return header, rows


def in_sync(csv_path: Path) -> bool:
    """True if the .xlsx exists and holds exactly the CSV's header and rows (a drop-down workbook is not byte-reproducible, so content is what is compared)."""
    x = xlsx_path(csv_path)
    if not x.exists():
        return False
    header, rows = read_csv(csv_path)
    xh, xr = read_xlsx(x)
    return xh == header and [[r.get(h, '') for h in header] for r in xr] == [r + [''] * (len(header) - len(r)) for r in rows]


def main():
    for p in SPECS:
        if in_sync(p):
            print('unchanged', xlsx_path(p).relative_to(ROOT))
        else:
            print('wrote', write_xlsx(p).relative_to(ROOT))
    return 0


if __name__ == '__main__':
    sys.exit(main())
